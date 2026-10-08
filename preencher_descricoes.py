"""Preenche um relatório HTML editável a partir de planilha XLSX/XLSM.

Uso: python preencher_descricoes.py
Requer: beautifulsoup4 e openpyxl.
"""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook


def normalizar(texto):
    s = unicodedata.normalize("NFKD", str(texto or ""))
    return re.sub(r"[^a-z0-9]+", "", "".join(ch for ch in s if not unicodedata.combining(ch)).lower())


def id_elemento(valor):
    if valor is None or str(valor).strip() == "":
        return ""
    s = str(valor).strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return s


# Nomes aceitos nos cabeçalhos da planilha. Ajuste conforme necessário.
COLUNAS = {
    "id1": ("Element ID Item 1", "ElementID Item 1", "Item 1 Element ID", "ID Item 1", "ElementID1"),
    "id2": ("Element ID Item 2", "ElementID Item 2", "Item 2 Element ID", "ID Item 2", "ElementID2"),
    "descricao": ("Descricao", "Descrição", "Description"),
    "responsavel": ("Quem altera", "Responsavel", "Responsável"),
}


def localizar_cabecalho(linha):
    nomes = [normalizar(valor) for valor in linha]
    indices = {}
    for campo, alternativas in COLUNAS.items():
        validos = {normalizar(nome) for nome in alternativas}
        correspondencias = [i for i, nome in enumerate(nomes) if nome in validos]
        if len(correspondencias) != 1:
            return None
        indices[campo] = correspondencias[0]
    return indices


def carregar_planilha(caminho: Path):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    registros = defaultdict(set)
    linhas = 0
    cabecalhos = 0
    try:
        for aba in wb.worksheets:
            cabecalho = None
            for linha in aba.iter_rows(values_only=True):
                if cabecalho is None:
                    encontrado = localizar_cabecalho(linha)
                    if encontrado:
                        cabecalho = encontrado
                        cabecalhos += 1
                    continue
                def obter(campo):
                    i = cabecalho[campo]
                    return linha[i] if i < len(linha) else None
                a, b = id_elemento(obter("id1")), id_elemento(obter("id2"))
                if not a or not b:
                    continue
                chave = tuple(sorted((a, b)))
                descricao = str(obter("descricao") or "").strip()
                responsavel = str(obter("responsavel") or "").strip()
                registros[chave].add((descricao, responsavel))
                linhas += 1
    finally:
        wb.close()
    if cabecalhos == 0:
        raise ValueError(
            "Cabeçalhos não identificados na planilha. São necessárias as colunas "
            "'Element ID Item 1', 'Element ID Item 2', 'Descricao' e 'Quem altera'. "
            "Confira os nomes em COLUNAS no início do script."
        )
    unicos = {chave: next(iter(valores)) for chave, valores in registros.items() if len(valores) == 1}
    ambiguos = {chave for chave, valores in registros.items() if len(valores) > 1}
    return unicos, ambiguos, linhas


def nome_valor(par: Tag):
    return (
        par.find("span", class_="name", recursive=False),
        par.find("span", class_="value", recursive=False)
    )


def buscar_par(raiz: Tag, nome: str):
    for par in raiz.find_all("span", class_="namevaluepair", recursive=False):
        label, value = nome_valor(par)
        if label and value and normalizar(label.get_text(" ", strip=True)) == normalizar(nome):
            return par, value
    return None, None


def ids_do_viewpoint(viewpoint: Tag):
    ids = {}
    for cabecalho in viewpoint.find_all("h4", class_="clashobject"):
        titulo = normalizar(cabecalho.get_text(" ", strip=True))
        if titulo not in ("item1", "item2"):
            continue
        for node in cabecalho.next_siblings:
            if isinstance(node, Tag) and node.name == "h4" and "clashobject" in node.get("class", []):
                break
            if not isinstance(node, Tag) or "namevaluepair" not in node.get("class", []):
                continue
            label, value = nome_valor(node)
            if label and value and normalizar(label.get_text(" ", strip=True)) == "elementid":
                ids[titulo] = id_elemento(value.get_text(" ", strip=True))
                break
    return ids


def preencher_valor(soup: BeautifulSoup, par: Tag, valor: Tag, texto: str, *, textarea: bool):
    controle = valor.find("textarea" if textarea else "input")
    if controle is None:
        valor.clear()
        if textarea:
            controle = soup.new_tag("textarea", attrs={"class": "editavel descricao", "rows": "2"})
        else:
            controle = soup.new_tag("input", attrs={
                "type": "text", "maxlength": "3", "size": "3", "class": "editavel curto"
            })
        valor.append(controle)
    if textarea:
        controle.string = texto
    else:
        controle["value"] = texto[:3]


def atualizar_html(html: Path, planilha: Path, saida: Path):
    if html.resolve() == saida.resolve():
        raise ValueError("Escolha uma saída diferente do HTML original.")
    unicos, ambiguos, linhas = carregar_planilha(planilha)
    soup = BeautifulSoup(html.read_text(encoding="utf-8-sig"), "html.parser")
    encontrados = sem_correspondencia = ambiguidades = sem_ids = 0
    for viewpoint in soup.select("div.viewpoint"):
        ids = ids_do_viewpoint(viewpoint)
        if not ids.get("item1") or not ids.get("item2"):
            sem_ids += 1
            continue
        chave = tuple(sorted((ids["item1"], ids["item2"])))
        if chave in ambiguos:
            ambiguidades += 1
            continue
        if chave not in unicos:
            sem_correspondencia += 1
            continue
        descricao, responsavel = unicos[chave]
        _, campo_desc = buscar_par(viewpoint, "Description")
        _, campo_resp = buscar_par(viewpoint, "Responsável")
        if campo_desc is None or campo_resp is None:
            sem_ids += 1
            continue
        # Valores vazios na planilha não apagam dados preenchidos manualmente.
        if descricao:
            preencher_valor(soup, viewpoint, campo_desc, descricao, textarea=True)
        if responsavel:
            preencher_valor(soup, viewpoint, campo_resp, responsavel, textarea=False)
        encontrados += 1
    saida.parent.mkdir(parents=True, exist_ok=True)
    saida.write_text(str(soup), encoding="utf-8")
    return dict(atualizados=encontrados, sem_correspondencia=sem_correspondencia,
                ambiguos=ambiguidades, sem_ids_ou_campos=sem_ids, linhas_planilha=linhas)


class Interface(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("EditHTML | Preencher descrições da planilha")
        self.geometry("790x380")
        self.minsize(650, 340)
        self.html = tk.StringVar()
        self.planilha = tk.StringVar()
        self.saida = tk.StringVar()
        self.status = tk.StringVar(value="Selecione o HTML editável, a planilha e o destino.")
        quadro = ttk.Frame(self, padding=22)
        quadro.pack(fill="both", expand=True)
        ttk.Label(quadro, text="Preencher descrições e responsáveis", font=("Segoe UI", 16, "bold")).pack(anchor="w", pady=(0, 8))
        ttk.Label(quadro, text="Correspondência pelo par de Element IDs de Item 1 e Item 2.").pack(anchor="w", pady=(0, 17))
        self.linha(quadro, "HTML editável", self.html, self.escolher_html)
        self.linha(quadro, "Planilha XLSX/XLSM", self.planilha, self.escolher_planilha)
        self.linha(quadro, "Salvar HTML em", self.saida, self.escolher_saida)
        ttk.Label(quadro, textvariable=self.status, wraplength=700).pack(anchor="w", pady=(18, 10))
        ttk.Button(quadro, text="Preencher e salvar HTML", command=self.executar).pack(anchor="e")

    def linha(self, quadro, nome, variavel, acao):
        linha = ttk.Frame(quadro)
        linha.pack(fill="x", pady=6)
        ttk.Label(linha, text=nome, width=20).pack(side="left")
        ttk.Entry(linha, textvariable=variavel).pack(side="left", fill="x", expand=True, padx=8)
        ttk.Button(linha, text="Selecionar...", command=acao).pack(side="right")

    def escolher_html(self):
        p = filedialog.askopenfilename(filetypes=[("HTML", "*.html *.htm")])
        if p:
            self.html.set(p)
            if not self.saida.get():
                origem = Path(p)
                self.saida.set(str(origem.with_name(origem.stem + "_preenchido.html")))

    def escolher_planilha(self):
        p = filedialog.askopenfilename(filetypes=[("Planilhas Excel", "*.xlsx *.xlsm")])
        if p:
            self.planilha.set(p)

    def escolher_saida(self):
        atual = Path(self.saida.get() or "relatorio_preenchido.html")
        p = filedialog.asksaveasfilename(defaultextension=".html", initialfile=atual.name,
                                         initialdir=str(atual.parent), filetypes=[("HTML", "*.html")])
        if p:
            self.saida.set(p)

    def executar(self):
        if not all(v.get().strip() for v in (self.html, self.planilha, self.saida)):
            messagebox.showwarning("Campos obrigatórios", "Informe os três caminhos.")
            return
        try:
            resultado = atualizar_html(Path(self.html.get()), Path(self.planilha.get()), Path(self.saida.get()))
        except Exception as exc:
            messagebox.showerror("Erro", str(exc))
            self.status.set("Falha no processamento.")
            return
        mensagem = (
            f"HTML salvo em: {self.saida.get()}\n"
            f"Interferências atualizadas: {resultado['atualizados']}\n"
            f"Sem correspondência: {resultado['sem_correspondencia']}\n"
            f"Com dados conflitantes na planilha: {resultado['ambiguos']}\n"
            f"Sem IDs ou campos necessários: {resultado['sem_ids_ou_campos']}"
        )
        self.status.set(mensagem)
        messagebox.showinfo("Concluído", mensagem)


if __name__ == "__main__":
    Interface().mainloop()
